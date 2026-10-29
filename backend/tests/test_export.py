import io
import pytest
from httpx import AsyncClient, ASGITransport
import openpyxl

from api.main import app


@pytest.mark.asyncio
async def test_export_csv_format():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        res = await ac.get("/api/licitaciones/export?format=csv&limit=10")
        assert res.status_code == 200
        assert "text/csv" in res.headers["content-type"]
        assert "attachment; filename=" in res.headers.get("content-disposition", "")
        assert res.headers["content-disposition"].endswith('.csv"')

        # Verificar presencia de BOM de UTF-8 (\xef\xbb\xbf)
        content_bytes = res.content
        assert content_bytes.startswith(b"\xef\xbb\xbf")

        # Decodificar y verificar delimitador ';' y cabeceras
        text = content_bytes.decode("utf-8-sig")
        lines = text.split("\r\n")
        assert len(lines) >= 2
        header_line = lines[0]
        assert ";" in header_line
        assert "OCID" in header_line
        assert "Rubro (IA)" in header_line
        assert "Descripción Limpia (IA)" in header_line


@pytest.mark.asyncio
async def test_export_xlsx_format():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        res = await ac.get("/api/licitaciones/export?format=xlsx&limit=10")
        assert res.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.headers["content-type"]
        assert "attachment; filename=" in res.headers.get("content-disposition", "")
        assert res.headers["content-disposition"].endswith('.xlsx"')

        # Verificar firma de archivo ZIP / XLSX (PK\x03\x04)
        assert res.content.startswith(b"PK\x03\x04")

        # Cargar con openpyxl para verificar validez y estructura
        bio = io.BytesIO(res.content)
        wb = openpyxl.load_workbook(bio)
        assert "Licitaciones SEACE" in wb.sheetnames
        ws = wb["Licitaciones SEACE"]
        assert ws.cell(row=1, column=1).value == "OCID"
        assert ws.cell(row=1, column=4).value == "Rubro (IA)"


@pytest.mark.asyncio
async def test_export_invalid_format():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        res = await ac.get("/api/licitaciones/export?format=pdf")
        assert res.status_code == 400
        assert "Formato no soportado" in res.json().get("detail", "")
